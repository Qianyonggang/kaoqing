from __future__ import annotations

import os
import sqlite3
from datetime import date, datetime
from functools import wraps
from io import BytesIO

from flask import (
    Flask,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from werkzeug.security import check_password_hash, generate_password_hash
from openpyxl import Workbook


# 创建 Flask 应用
app = Flask(__name__)
app.secret_key = os.environ.get("APP_SECRET", "dev-secret-key")
DB_PATH = os.environ.get("APP_DB", os.path.join(os.path.dirname(__file__), "app.db"))


def get_db() -> sqlite3.Connection:
    """获取数据库连接，确保行以字典形式返回。"""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    """初始化数据库表结构。"""
    with get_db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS companies (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                company_id INTEGER NOT NULL,
                username TEXT NOT NULL,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL,
                can_manage_admins INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                UNIQUE(company_id, username)
            );

            CREATE TABLE IF NOT EXISTS teams (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                company_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                created_by INTEGER NOT NULL,
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS employees (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                company_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                phone TEXT,
                daily_wage REAL NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS team_members (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                team_id INTEGER NOT NULL,
                employee_id INTEGER NOT NULL,
                UNIQUE(team_id, employee_id)
            );

            CREATE TABLE IF NOT EXISTS attendance (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                company_id INTEGER NOT NULL,
                team_id INTEGER NOT NULL,
                employee_id INTEGER NOT NULL,
                work_date TEXT NOT NULL,
                work_units REAL NOT NULL,
                notes TEXT,
                created_by INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                UNIQUE(team_id, employee_id, work_date)
            );

            CREATE TABLE IF NOT EXISTS advances (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                company_id INTEGER NOT NULL,
                employee_id INTEGER NOT NULL,
                advance_date TEXT NOT NULL,
                amount REAL NOT NULL,
                notes TEXT,
                created_by INTEGER NOT NULL,
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                company_id INTEGER NOT NULL,
                user_id INTEGER NOT NULL,
                action TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            """
        )


@app.before_request
def ensure_db() -> None:
    """每次请求前确保数据库已初始化。"""
    init_db()


def log_action(company_id: int, user_id: int, action: str) -> None:
    """记录操作日志，供公司创建者查看。"""
    with get_db() as conn:
        conn.execute(
            "INSERT INTO logs (company_id, user_id, action, created_at) VALUES (?, ?, ?, ?)",
            (company_id, user_id, action, datetime.now().isoformat(timespec="seconds")),
        )


def login_required(view):
    """登录保护装饰器。"""

    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login"))
        return view(*args, **kwargs)

    return wrapped


def owner_required(view):
    """公司创建者权限装饰器。"""

    @wraps(view)
    def wrapped(*args, **kwargs):
        if session.get("role") != "owner":
            flash("需要公司创建者权限")
            return redirect(url_for("dashboard"))
        return view(*args, **kwargs)

    return wrapped


@app.route("/")
@login_required
def dashboard():
    """系统首页，展示公司信息与快速入口。"""
    company_id = session["company_id"]
    with get_db() as conn:
        company = conn.execute("SELECT * FROM companies WHERE id = ?", (company_id,)).fetchone()
        team_count = conn.execute(
            "SELECT COUNT(*) AS total FROM teams WHERE company_id = ?", (company_id,)
        ).fetchone()["total"]
        employee_count = conn.execute(
            "SELECT COUNT(*) AS total FROM employees WHERE company_id = ?", (company_id,)
        ).fetchone()["total"]
    return render_template(
        "dashboard.html",
        company=company,
        team_count=team_count,
        employee_count=employee_count,
    )


@app.route("/register", methods=["GET", "POST"])
def register():
    """注册公司账户并创建公司管理员(创建者)。"""
    if request.method == "POST":
        company_name = request.form.get("company_name", "").strip()
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        if not company_name or not username or not password:
            flash("请填写完整信息")
            return render_template("auth/register.html")
        with get_db() as conn:
            cursor = conn.execute(
                "INSERT INTO companies (name, created_at) VALUES (?, ?)",
                (company_name, datetime.now().isoformat(timespec="seconds")),
            )
            company_id = cursor.lastrowid
            conn.execute(
                """
                INSERT INTO users (company_id, username, password_hash, role, can_manage_admins, created_at)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    company_id,
                    username,
                    generate_password_hash(password),
                    "owner",
                    1,
                    datetime.now().isoformat(timespec="seconds"),
                ),
            )
        flash("注册成功，请登录")
        return redirect(url_for("login"))
    return render_template("auth/register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    """登录公司账户。"""
    if request.method == "POST":
        company_name = request.form.get("company_name", "").strip()
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        if not company_name or not username or not password:
            flash("请输入公司名称、账号和密码")
            return render_template("auth/login.html")
        with get_db() as conn:
            user = conn.execute(
                """
                SELECT users.* FROM users
                JOIN companies ON users.company_id = companies.id
                WHERE users.username = ? AND companies.name = ?
                """,
                (username, company_name),
            ).fetchone()
        if not user or not check_password_hash(user["password_hash"], password):
            flash("账号或密码错误")
            return render_template("auth/login.html")
        session.clear()
        session["user_id"] = user["id"]
        session["company_id"] = user["company_id"]
        session["role"] = user["role"]
        session["can_manage_admins"] = bool(user["can_manage_admins"])
        log_action(user["company_id"], user["id"], "用户登录")
        return redirect(url_for("dashboard"))
    return render_template("auth/login.html")


@app.route("/logout")
@login_required
def logout():
    """退出登录。"""
    log_action(session["company_id"], session["user_id"], "用户退出登录")
    session.clear()
    return redirect(url_for("login"))


@app.route("/teams", methods=["GET", "POST"])
@login_required
def teams():
    """团队列表与创建。"""
    company_id = session["company_id"]
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        if not name:
            flash("团队名称不能为空")
        else:
            with get_db() as conn:
                conn.execute(
                    """
                    INSERT INTO teams (company_id, name, created_by, created_at)
                    VALUES (?, ?, ?, ?)
                    """,
                    (company_id, name, session["user_id"], datetime.now().isoformat(timespec="seconds")),
                )
            log_action(company_id, session["user_id"], f"创建团队：{name}")
            flash("团队创建成功")
            return redirect(url_for("teams"))
    with get_db() as conn:
        team_list = conn.execute(
            """
            SELECT teams.*, users.username AS creator
            FROM teams
            JOIN users ON teams.created_by = users.id
            WHERE teams.company_id = ?
            ORDER BY teams.id DESC
            """,
            (company_id,),
        ).fetchall()
    return render_template("teams/teams.html", teams=team_list)


@app.route("/teams/<int:team_id>", methods=["GET", "POST"])
@login_required
def team_detail(team_id: int):
    """团队详情与成员管理。"""
    company_id = session["company_id"]
    with get_db() as conn:
        team = conn.execute(
            "SELECT * FROM teams WHERE id = ? AND company_id = ?", (team_id, company_id)
        ).fetchone()
        if not team:
            flash("团队不存在")
            return redirect(url_for("teams"))
        search_keyword = request.args.get("q", "").strip()
        employee_query = "SELECT * FROM employees WHERE company_id = ?"
        employee_params: list[object] = [company_id]
        if search_keyword:
            employee_query += " AND (name LIKE ? OR phone LIKE ?)"
            employee_params.extend([f"%{search_keyword}%", f"%{search_keyword}%"])
        employee_query += " ORDER BY id DESC"
        employees = conn.execute(employee_query, employee_params).fetchall()
        members = conn.execute(
            """
            SELECT employees.*
            FROM team_members
            JOIN employees ON team_members.employee_id = employees.id
            WHERE team_members.team_id = ?
            """,
            (team_id,),
        ).fetchall()
    if request.method == "POST":
        action = request.form.get("action")
        if action == "add_existing":
            employee_id = request.form.get("employee_id")
            if not employee_id:
                flash("请选择员工")
            else:
                with get_db() as conn:
                    conn.execute(
                        "INSERT OR IGNORE INTO team_members (team_id, employee_id) VALUES (?, ?)",
                        (team_id, employee_id),
                    )
                log_action(company_id, session["user_id"], f"团队{team['name']}新增员工ID:{employee_id}")
                flash("成员添加成功")
            return redirect(url_for("team_detail", team_id=team_id, q=search_keyword))
        if action == "create_employee":
            name = request.form.get("name", "").strip()
            phone = request.form.get("phone", "").strip()
            daily_wage = request.form.get("daily_wage", "0").strip() or "0"
            if not name:
                flash("员工姓名不能为空")
            else:
                with get_db() as conn:
                    cursor = conn.execute(
                        """
                        INSERT INTO employees (company_id, name, phone, daily_wage, created_at)
                        VALUES (?, ?, ?, ?, ?)
                        """,
                        (
                            company_id,
                            name,
                            phone,
                            float(daily_wage),
                            datetime.now().isoformat(timespec="seconds"),
                        ),
                    )
                    conn.execute(
                        "INSERT OR IGNORE INTO team_members (team_id, employee_id) VALUES (?, ?)",
                        (team_id, cursor.lastrowid),
                    )
                log_action(company_id, session["user_id"], f"团队{team['name']}新增员工：{name}")
                flash("员工已创建并加入团队")
            return redirect(url_for("team_detail", team_id=team_id, q=search_keyword))
    return render_template(
        "teams/team_detail.html",
        team=team,
        employees=employees,
        members=members,
        search_keyword=search_keyword,
    )


@app.route("/teams/<int:team_id>/attendance", methods=["GET", "POST"])
@login_required
def team_attendance(team_id: int):
    """团队每日考勤登记。"""
    company_id = session["company_id"]
    work_date = request.args.get("date") or date.today().isoformat()
    search_keyword = request.args.get("q", "").strip()
    today = date.today()
    try:
        selected_date = date.fromisoformat(work_date)
    except ValueError:
        selected_date = today
    if selected_date > today:
        flash("不能登记未来日期的考勤")
        selected_date = today
    work_date = selected_date.isoformat()
    month_prefix = selected_date.strftime("%Y-%m")
    with get_db() as conn:
        team = conn.execute(
            "SELECT * FROM teams WHERE id = ? AND company_id = ?", (team_id, company_id)
        ).fetchone()
        if not team:
            flash("团队不存在")
            return redirect(url_for("teams"))
        member_query = """
            SELECT employees.*
            FROM team_members
            JOIN employees ON team_members.employee_id = employees.id
            WHERE team_members.team_id = ?
        """
        member_params: list[object] = [team_id]
        if search_keyword:
            member_query += " AND (employees.name LIKE ? OR employees.phone LIKE ?)"
            member_params.extend([f"%{search_keyword}%", f"%{search_keyword}%"])
        member_query += " ORDER BY employees.id DESC"
        members = conn.execute(member_query, member_params).fetchall()
        attendance_map = {
            row["employee_id"]: row
            for row in conn.execute(
                """
                SELECT * FROM attendance
                WHERE team_id = ? AND work_date = ?
                """,
                (team_id, work_date),
            ).fetchall()
        }
        total_rows = conn.execute(
            """
            SELECT employee_id, SUM(work_units) AS total_units
            FROM attendance
            WHERE team_id = ? AND work_date LIKE ?
            GROUP BY employee_id
            """,
            (team_id, f"{month_prefix}%"),
        ).fetchall()
        total_units_map = {row["employee_id"]: row["total_units"] or 0 for row in total_rows}
    if request.method == "POST":
        if selected_date > today:
            flash("不能登记未来日期的考勤")
            return redirect(url_for("team_attendance", team_id=team_id, date=today.isoformat()))
        with get_db() as conn:
            for member in members:
                work_units = request.form.get(f"work_units_{member['id']}", "0").strip() or "0"
                notes = request.form.get(f"notes_{member['id']}", "").strip()
                try:
                    units_value = float(work_units)
                except ValueError:
                    units_value = 0
                if units_value > 1:
                    units_value = 1
                if units_value < -1:
                    units_value = -1
                conn.execute(
                    "DELETE FROM attendance WHERE team_id = ? AND employee_id = ? AND work_date = ?",
                    (team_id, member["id"], work_date),
                )
                conn.execute(
                    """
                    INSERT INTO attendance (
                        company_id, team_id, employee_id, work_date, work_units, notes, created_by, created_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        company_id,
                        team_id,
                        member["id"],
                        work_date,
                        units_value,
                        notes,
                        session["user_id"],
                        datetime.now().isoformat(timespec="seconds"),
                    ),
                )
        log_action(company_id, session["user_id"], f"更新团队{team['name']} {work_date}考勤")
        flash("考勤已保存")
        return redirect(url_for("team_attendance", team_id=team_id, date=work_date))
    return render_template(
        "teams/attendance.html",
        team=team,
        members=members,
        attendance_map=attendance_map,
        work_date=work_date,
        search_keyword=search_keyword,
        total_units_map=total_units_map,
        month_prefix=month_prefix,
        today=today.isoformat(),
    )


@app.route("/employees", methods=["GET", "POST"])
@login_required
def employees():
    """员工列表与创建。"""
    company_id = session["company_id"]
    search_keyword = request.args.get("q", "").strip()
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        phone = request.form.get("phone", "").strip()
        daily_wage = request.form.get("daily_wage", "0").strip() or "0"
        if not name:
            flash("员工姓名不能为空")
        else:
            with get_db() as conn:
                conn.execute(
                    """
                    INSERT INTO employees (company_id, name, phone, daily_wage, created_at)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (
                        company_id,
                        name,
                        phone,
                        float(daily_wage),
                        datetime.now().isoformat(timespec="seconds"),
                    ),
                )
            log_action(company_id, session["user_id"], f"新增员工：{name}")
            flash("员工创建成功")
            return redirect(url_for("employees", q=search_keyword))
    with get_db() as conn:
        employee_query = "SELECT * FROM employees WHERE company_id = ?"
        employee_params: list[object] = [company_id]
        if search_keyword:
            employee_query += " AND (name LIKE ? OR phone LIKE ?)"
            employee_params.extend([f"%{search_keyword}%", f"%{search_keyword}%"])
        employee_query += " ORDER BY id DESC"
        employee_list = conn.execute(employee_query, employee_params).fetchall()
    return render_template(
        "employees/employees.html",
        employees=employee_list,
        search_keyword=search_keyword,
    )


@app.route("/employees/<int:employee_id>/delete", methods=["POST"])
@login_required
def delete_employee(employee_id: int):
    """删除员工及相关数据。"""
    company_id = session["company_id"]
    with get_db() as conn:
        employee = conn.execute(
            "SELECT * FROM employees WHERE id = ? AND company_id = ?",
            (employee_id, company_id),
        ).fetchone()
        if not employee:
            flash("员工不存在")
            return redirect(url_for("employees"))
        conn.execute("DELETE FROM attendance WHERE employee_id = ?", (employee_id,))
        conn.execute("DELETE FROM advances WHERE employee_id = ?", (employee_id,))
        conn.execute("DELETE FROM team_members WHERE employee_id = ?", (employee_id,))
        conn.execute("DELETE FROM employees WHERE id = ?", (employee_id,))
    log_action(company_id, session["user_id"], f"删除员工：{employee['name']}")
    flash("员工已删除")
    return redirect(url_for("employees"))


@app.route("/employees/<int:employee_id>")
@login_required
def employee_detail(employee_id: int):
    """员工详情与考勤汇总。"""
    company_id = session["company_id"]
    month = request.args.get("month") or datetime.now().strftime("%Y-%m")
    with get_db() as conn:
        employee = conn.execute(
            "SELECT * FROM employees WHERE id = ? AND company_id = ?",
            (employee_id, company_id),
        ).fetchone()
        if not employee:
            flash("员工不存在")
            return redirect(url_for("employees"))
        attendance_rows = conn.execute(
            """
            SELECT attendance.*, teams.name AS team_name, users.username AS admin_name
            FROM attendance
            JOIN teams ON attendance.team_id = teams.id
            JOIN users ON attendance.created_by = users.id
            WHERE attendance.employee_id = ? AND attendance.work_date LIKE ?
            ORDER BY attendance.work_date DESC
            """,
            (employee_id, f"{month}%"),
        ).fetchall()
        advances = conn.execute(
            """
            SELECT * FROM advances
            WHERE employee_id = ? AND advance_date LIKE ?
            ORDER BY advance_date DESC
            """,
            (employee_id, f"{month}%"),
        ).fetchall()
    total_units = sum(row["work_units"] for row in attendance_rows)
    total_wage = total_units * employee["daily_wage"]
    total_advance = sum(row["amount"] for row in advances)
    remaining = total_wage - total_advance
    return render_template(
        "employees/employee_detail.html",
        employee=employee,
        attendance_rows=attendance_rows,
        advances=advances,
        month=month,
        total_units=total_units,
        total_wage=total_wage,
        total_advance=total_advance,
        remaining=remaining,
    )


@app.route("/employees/<int:employee_id>/advances", methods=["GET", "POST"])
@login_required
def employee_advances(employee_id: int):
    """记录员工借支。"""
    company_id = session["company_id"]
    with get_db() as conn:
        employee = conn.execute(
            "SELECT * FROM employees WHERE id = ? AND company_id = ?",
            (employee_id, company_id),
        ).fetchone()
        if not employee:
            flash("员工不存在")
            return redirect(url_for("employees"))
    if request.method == "POST":
        amount = request.form.get("amount", "0").strip() or "0"
        advance_date = request.form.get("advance_date") or date.today().isoformat()
        notes = request.form.get("notes", "").strip()
        with get_db() as conn:
            conn.execute(
                """
                INSERT INTO advances (company_id, employee_id, advance_date, amount, notes, created_by, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    company_id,
                    employee_id,
                    advance_date,
                    float(amount),
                    notes,
                    session["user_id"],
                    datetime.now().isoformat(timespec="seconds"),
                ),
            )
        log_action(company_id, session["user_id"], f"记录员工{employee['name']}借支 {amount}")
        flash("借支记录已保存")
        return redirect(url_for("employee_advances", employee_id=employee_id))
    with get_db() as conn:
        advances = conn.execute(
            """
            SELECT * FROM advances
            WHERE employee_id = ?
            ORDER BY advance_date DESC
            """,
            (employee_id,),
        ).fetchall()
    return render_template(
        "employees/advances.html", employee=employee, advances=advances
    )


@app.route("/admin/users", methods=["GET", "POST"])
@login_required
@owner_required
def admin_users():
    """公司创建者管理管理员账号。"""
    company_id = session["company_id"]
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        can_manage_admins = 1 if request.form.get("can_manage_admins") else 0
        if not username or not password:
            flash("请输入账号和密码")
        else:
            try:
                with get_db() as conn:
                    conn.execute(
                        """
                        INSERT INTO users (company_id, username, password_hash, role, can_manage_admins, created_at)
                        VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (
                            company_id,
                            username,
                            generate_password_hash(password),
                            "admin",
                            can_manage_admins,
                            datetime.now().isoformat(timespec="seconds"),
                        ),
                    )
                log_action(company_id, session["user_id"], f"新增管理员：{username}")
                flash("管理员创建成功")
                return redirect(url_for("admin_users"))
            except sqlite3.IntegrityError:
                flash("账号已存在")
    with get_db() as conn:
        admins = conn.execute(
            "SELECT * FROM users WHERE company_id = ? ORDER BY id DESC", (company_id,)
        ).fetchall()
    return render_template("admin/users.html", admins=admins)


@app.route("/admin/users/<int:user_id>/toggle", methods=["POST"])
@login_required
@owner_required
def toggle_admin_permission(user_id: int):
    """切换管理员的授权管理权限。"""
    company_id = session["company_id"]
    with get_db() as conn:
        user = conn.execute(
            "SELECT * FROM users WHERE id = ? AND company_id = ?", (user_id, company_id)
        ).fetchone()
        if not user or user["role"] == "owner":
            flash("无法修改该账号")
            return redirect(url_for("admin_users"))
        new_value = 0 if user["can_manage_admins"] else 1
        conn.execute(
            "UPDATE users SET can_manage_admins = ? WHERE id = ?", (new_value, user_id)
        )
    log_action(company_id, session["user_id"], f"修改管理员权限：{user['username']}")
    flash("权限已更新")
    return redirect(url_for("admin_users"))


@app.route("/logs")
@login_required
@owner_required
def logs():
    """查看操作日志。"""
    company_id = session["company_id"]
    with get_db() as conn:
        log_list = conn.execute(
            """
            SELECT logs.*, users.username
            FROM logs
            JOIN users ON logs.user_id = users.id
            WHERE logs.company_id = ?
            ORDER BY logs.id DESC
            LIMIT 200
            """,
            (company_id,),
        ).fetchall()
    return render_template("admin/logs.html", logs=log_list)


@app.route("/reports")
@login_required
def reports():
    """查看汇总报表。"""
    company_id = session["company_id"]
    month = request.args.get("month") or datetime.now().strftime("%Y-%m")
    team_ids = [int(team_id) for team_id in request.args.getlist("team_id") if team_id.isdigit()]
    with get_db() as conn:
        teams = conn.execute(
            "SELECT * FROM teams WHERE company_id = ? ORDER BY id DESC", (company_id,)
        ).fetchall()
        employees = conn.execute(
            "SELECT * FROM employees WHERE company_id = ? ORDER BY id DESC", (company_id,)
        ).fetchall()
        attendance_query = """
            SELECT employee_id, SUM(work_units) AS total_units
            FROM attendance
            WHERE company_id = ? AND work_date LIKE ?
        """
        attendance_params: list[object] = [company_id, f"{month}%"]
        if team_ids:
            placeholders = ",".join("?" for _ in team_ids)
            attendance_query += f" AND team_id IN ({placeholders})"
            attendance_params.extend(team_ids)
        attendance_query += " GROUP BY employee_id"
        attendance_rows = conn.execute(attendance_query, attendance_params).fetchall()
        advances_rows = conn.execute(
            """
            SELECT employee_id, SUM(amount) AS total_amount
            FROM advances
            WHERE company_id = ? AND advance_date LIKE ?
            GROUP BY employee_id
            """,
            (company_id, f"{month}%"),
        ).fetchall()
    attendance_map = {row["employee_id"]: row["total_units"] or 0 for row in attendance_rows}
    advance_map = {row["employee_id"]: row["total_amount"] or 0 for row in advances_rows}
    report_data = []
    for employee in employees:
        total_units = attendance_map.get(employee["id"], 0)
        total_wage = total_units * employee["daily_wage"]
        total_advance = advance_map.get(employee["id"], 0)
        remaining = total_wage - total_advance
        report_data.append(
            {
                "employee": employee,
                "total_units": total_units,
                "total_wage": total_wage,
                "total_advance": total_advance,
                "remaining": remaining,
            }
        )
    return render_template(
        "reports/reports.html",
        month=month,
        report_data=report_data,
        teams=teams,
        selected_team_ids=team_ids,
    )


@app.route("/reports/export")
@login_required
@owner_required
def export_report():
    """导出 Excel 报表。"""
    company_id = session["company_id"]
    month = request.args.get("month") or datetime.now().strftime("%Y-%m")
    year = month.split("-")[0]
    team_ids = [int(team_id) for team_id in request.args.getlist("team_id") if team_id.isdigit()]
    with get_db() as conn:
        team_query = "SELECT * FROM teams WHERE company_id = ?"
        team_params: list[object] = [company_id]
        if team_ids:
            placeholders = ",".join("?" for _ in team_ids)
            team_query += f" AND id IN ({placeholders})"
            team_params.extend(team_ids)
        team_query += " ORDER BY id DESC"
        teams = conn.execute(team_query, team_params).fetchall()
        employees = conn.execute(
            "SELECT * FROM employees WHERE company_id = ?", (company_id,)
        ).fetchall()
        attendance_query = """
            SELECT * FROM attendance
            WHERE company_id = ? AND work_date LIKE ?
        """
        attendance_params: list[object] = [company_id, f"{year}-%"]
        if team_ids:
            placeholders = ",".join("?" for _ in team_ids)
            attendance_query += f" AND team_id IN ({placeholders})"
            attendance_params.extend(team_ids)
        attendance_rows = conn.execute(attendance_query, attendance_params).fetchall()
        advances_rows = conn.execute(
            """
            SELECT * FROM advances
            WHERE company_id = ? AND advance_date LIKE ?
            """,
            (company_id, f"{year}-%"),
        ).fetchall()
        team_members = conn.execute(
            "SELECT * FROM team_members"
        ).fetchall()
    attendance_by_team = {}
    for row in attendance_rows:
        attendance_by_team.setdefault((row["team_id"], row["employee_id"]), 0)
        attendance_by_team[(row["team_id"], row["employee_id"])] += row["work_units"]
    advances_map = {}
    for row in advances_rows:
        advances_map.setdefault(row["employee_id"], 0)
        advances_map[row["employee_id"]] += row["amount"]
    employee_map = {emp["id"]: emp for emp in employees}
    members_by_team = {}
    for member in team_members:
        members_by_team.setdefault(member["team_id"], []).append(member["employee_id"])
    team_name_map = {team["id"]: team["name"] for team in teams}

    monthly_labels = [f"{idx}月" for idx in range(1, 13)]
    monthly_totals = {emp["id"]: {idx: 0 for idx in range(1, 13)} for emp in employees}
    monthly_by_team: dict[tuple[int, int], dict[int, float]] = {}
    for row in attendance_rows:
        month_index = int(row["work_date"].split("-")[1])
        monthly_totals[row["employee_id"]][month_index] += row["work_units"]
        key = (row["team_id"], row["employee_id"])
        monthly_by_team.setdefault(key, {idx: 0 for idx in range(1, 13)})
        monthly_by_team[key][month_index] += row["work_units"]

    wb = Workbook()
    summary_sheet = wb.active
    summary_sheet.title = "汇总"
    summary_sheet.append([
        "员工姓名",
        "联系方式",
        "所属团队",
        *monthly_labels,
        "总天数",
        "单日工资",
        "借支",
        "总工资",
        "剩余工资",
    ])
    for employee in employees:
        total_units = sum(monthly_totals[employee["id"]].values())
        total_wage = total_units * employee["daily_wage"]
        total_advance = advances_map.get(employee["id"], 0)
        remaining = total_wage - total_advance
        team_names = sorted({
            team_name_map.get(team_id, "")
            for (team_id, emp_id) in attendance_by_team
            if emp_id == employee["id"]
        })
        summary_sheet.append(
            [
                employee["name"],
                employee["phone"],
                "、".join([name for name in team_names if name]) or "-",
                *[monthly_totals[employee["id"]][idx] for idx in range(1, 13)],
                total_units,
                employee["daily_wage"],
                total_advance,
                total_wage,
                remaining,
            ]
        )

    for team in teams:
        sheet = wb.create_sheet(title=team["name"][:31])
        sheet.append([
            "员工姓名",
            "联系方式",
            "所属团队",
            *monthly_labels,
            "总天数",
            "单日工资",
            "借支",
            "总工资",
            "剩余工资",
        ])
        for emp_id in members_by_team.get(team["id"], []):
            employee = employee_map.get(emp_id)
            if not employee:
                continue
            total_units = attendance_by_team.get((team["id"], emp_id), 0)
            total_wage = total_units * employee["daily_wage"]
            total_advance = advances_map.get(emp_id, 0)
            remaining = total_wage - total_advance
            monthly_data = monthly_by_team.get((team["id"], emp_id), {idx: 0 for idx in range(1, 13)})
            sheet.append(
                [
                    employee["name"],
                    employee["phone"],
                    team["name"],
                    *[monthly_data[idx] for idx in range(1, 13)],
                    total_units,
                    employee["daily_wage"],
                    total_advance,
                    total_wage,
                    remaining,
                ]
            )

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"attendance_report_{year}.xlsx"
    log_action(company_id, session["user_id"], f"导出报表 {year}")
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
